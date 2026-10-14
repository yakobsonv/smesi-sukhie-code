import json, re
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

R = json.load(open('recon/records.json')); P = {p['id']: p for p in json.load(open('recon/products.json'))}
S = json.load(open('recon/sections.json'))
SITE = {'perel': 'Perel', 'dauer': 'Dauer', 'osnovit': 'Основит'}
STALE = {'основит флайвэлл pc24/1 ml', 'основит стартвэлл pc21', 'основит стартвэлл pc21 m'}
T2S = {}
for d in S['sections']:
    for t in d['types']:
        T2S[t] = d['section'].replace(' (коротко + ссылка на сайт «Цветные смеси»)', '')
ND = {'primer': 'Грунтовки', 'paint': 'Краски и эмали', 'decorative_plaster_paste': 'Декоративные штукатурки готовые (пасты)',
      'putty_ready_paste': 'Шпаклёвки готовые (пасты)', 'additive': 'Добавки в растворы', 'hydrophobizer': 'Гидрофобизаторы',
      'efflorescence_remover': 'Средства от высолов', 'waterproofing_liquid': 'Гидроизоляция готовая (жидкая)',
      'tile_grout_epoxy': 'Затирки эпоксидные', 'cleaning_care': 'Бытовые средства ухода и очистки', 'sealant': 'Герметики',
      'wallpaper_glue': 'Клей для обоев', 'polymer_coating': 'Полимерные покрытия', 'other': 'Прочее'}

def section(r):
    if r['is_dry_mix'] in ('yes', 'unclear') and r['type'] in T2S:
        return T2S[r['type']]
    if r['type'] == 'paint' and r['is_dry_mix'] == 'yes':
        return 'Краска сухая (цементно-известковая)'
    return 'Сопутствующие, не сухие: ' + ND.get(r['type'], r['type']) if r['is_dry_mix'] == 'no' or r['type'] not in T2S else T2S[r['type']]

F = 'Arial'
hdr_font = Font(name=F, bold=True, color='FFFFFF'); hdr_fill = PatternFill('solid', fgColor='4A3F35')
body = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True)
thin = Side(style='thin', color='D9D4CC'); border = Border(bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')

def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=i, value=c); cell.font = hdr_font; cell.fill = hdr_fill; cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'; ws.row_dimensions[1].height = 32

wb = Workbook()
# --- Товары (series level)
wt = wb.active; wt.title = 'Товары'
cols = ['Производитель', 'Раздел будущего сайта', 'Сухая смесь', 'Товар (без цвета и фасовки)', 'Название на сайте производителя',
        'Карточек у производителя (цвета, фасовки)', 'Назначение — дословно', 'Фасовка — дословно', 'Расход — дословно',
        'Толщина слоя или шва — дословно', 'Норматив — дословно', 'Зимняя версия', 'Цвета — дословно', 'Ссылка на карточку', 'Примечание разбора']
header(wt, cols, [12, 34, 9, 30, 44, 12, 60, 22, 30, 22, 20, 14, 20, 50, 50])
WV = {'winter_only': 'зимняя', 'has_winter_variant': 'есть', 'no': 'нет', 'not_stated': '—'}
groups = defaultdict(list)
for r in R:
    if r['page_kind'] != 'product':
        continue
    site = P[r['id']]['site']
    k = r['series_key'].strip()
    if k.lower() in STALE:
        continue
    groups[(site, k)].append(r)
rows = []
for (site, k), v in groups.items():
    r0 = sorted(v, key=lambda r: (-len(r['purpose_quote']), r['id']))[0]
    pick = lambda f: next((x[f] for x in v if (x.get(f) or '').strip()), '')
    wvals = set(x.get('winter_version', 'not_stated') for x in v)
    wv = 'есть' if wvals & {'winter_only', 'has_winter_variant'} else ('нет' if wvals == {'no'} else '—')
    rows.append([SITE[site], section(r0), 'да' if any(x['is_dry_mix'] == 'yes' for x in v) else ('неясно' if any(x['is_dry_mix'] == 'unclear' for x in v) else 'нет'),
                 k, r0['name_exact'], len(v), r0['purpose_quote'], pick('packaging_quote'), pick('consumption_quote'),
                 pick('layer_or_joint_quote'), pick('normative_quote'), wv, pick('colors_quote'), P[r0['id']]['url'], r0['notes'][:400]])
order = {d['section'].replace(' (коротко + ссылка на сайт «Цветные смеси»)', ''): i for i, d in enumerate(S['sections'])}
rows.sort(key=lambda x: (x[2] != 'да', order.get(x[1], 99), x[1], x[0], x[3]))
for i, row in enumerate(rows, 2):
    for j, val in enumerate(row, 1):
        c = wt.cell(row=i, column=j, value=val); c.font = body; c.alignment = wrap; c.border = border
    wt.cell(row=i, column=14).hyperlink = row[13]
wt.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows)+1}"
NT = len(rows) + 1

# --- Сводка (formulas over Товары)
ws = wb.create_sheet('Сводка', 0)
ws['A1'] = 'Сухие строительные смеси — ассортимент трёх производителей (разведка 09.10.2026)'; ws['A1'].font = Font(name=F, size=13, bold=True)
ws['A2'] = 'Считаются товары без учёта цвета и фасовки. Источник — полный обход каталогов td-perel.ru, dauer.ru, osnovit.ru 09.10.2026 (лист «Как собрано»). Числа считаются формулами из листа «Товары».'; ws['A2'].font = Font(name=F, size=9, italic=True)
heads = ['Раздел будущего сайта', 'Perel', 'Dauer', 'Основит', 'Всего', 'У скольких производителей']
for i, h in enumerate(heads, 1):
    c = ws.cell(row=4, column=i, value=h); c.font = hdr_font; c.fill = hdr_fill
for col, w in zip('ABCDEF', [62, 9, 9, 10, 9, 16]):
    ws.column_dimensions[col].width = w
r = 5
secs = [d['section'].replace(' (коротко + ссылка на сайт «Цветные смеси»)', '') for d in S['sections']] + ['Краска сухая (цементно-известковая)']
for s in secs:
    ws.cell(row=r, column=1, value=s).font = body
    for j, m in enumerate(['Perel', 'Dauer', 'Основит'], 2):
        ws.cell(row=r, column=j, value=f'=COUNTIFS(Товары!$A$2:$A${NT},"{m}",Товары!$B$2:$B${NT},$A{r},Товары!$C$2:$C${NT},"<>нет")').font = body
    ws.cell(row=r, column=5, value=f'=SUM(B{r}:D{r})').font = bold
    ws.cell(row=r, column=6, value=f'=(B{r}>0)+(C{r}>0)+(D{r}>0)').font = body
    r += 1
ws.cell(row=r, column=1, value='Итого сухих смесей').font = bold
for j in range(2, 6):
    L = get_column_letter(j); ws.cell(row=r, column=j, value=f'=SUM({L}5:{L}{r-1})').font = bold
tot_row = r
r += 2
ws.cell(row=r, column=1, value='Сопутствующие не сухие товары (грунты, краски, пасты, добавки, бытовая химия)').font = body
for j, m in enumerate(['Perel', 'Dauer', 'Основит'], 2):
    ws.cell(row=r, column=j, value=f'=COUNTIFS(Товары!$A$2:$A${NT},"{m}",Товары!$C$2:$C${NT},"нет")').font = body
ws.cell(row=r, column=5, value=f'=SUM(B{r}:D{r})').font = bold
r += 1
ws.cell(row=r, column=1, value='Разделов, где позиции есть у всех трёх').font = body
ws.cell(row=r, column=5, value=f'=COUNTIF(F5:F{tot_row-1},3)').font = bold
r += 1
ws.cell(row=r, column=1, value='Разделов, где позиции есть только у одного').font = body
ws.cell(row=r, column=5, value=f'=COUNTIF(F5:F{tot_row-1},1)').font = bold
r += 2
ws.cell(row=r, column=1, value='Цветные кладочные растворы — тема соседнего сайта «Цветные кладочные смеси»: здесь коротко и ссылкой.').font = Font(name=F, size=9, italic=True)

# --- Карточки (all cards)
wc = wb.create_sheet('Карточки')
cc = ['Производитель', 'Тип страницы', 'Название (h1)', 'Товар', 'Вариант', 'Наш тип', 'Сухая', 'Назначение — дословно', 'Фасовка — дословно', 'Расход — дословно', 'Метки', 'Ссылка']
header(wc, cc, [11, 14, 46, 28, 22, 22, 8, 60, 22, 30, 16, 50])
for i, rr in enumerate(sorted(R, key=lambda x: (P[x['id']]['site'], x['series_key'], x['variant'])), 2):
    vals = [SITE[P[rr['id']]['site']], rr['page_kind'], rr['name_exact'], rr['series_key'], rr['variant'], rr['type'], rr['is_dry_mix'],
            rr['purpose_quote'], rr.get('packaging_quote', ''), rr.get('consumption_quote', ''), ', '.join(rr.get('status_flags') or []), P[rr['id']]['url']]
    for j, val in enumerate(vals, 1):
        c = wc.cell(row=i, column=j, value=val); c.font = body; c.alignment = wrap
wc.auto_filter.ref = f"A1:L{len(R)+1}"

# --- Как собрано
wm = wb.create_sheet('Как собрано')
notes = [
    'Как собрано',
    '09.10.2026 обойдены все страницы раздела /catalog/ трёх сайтов: td-perel.ru (249 страниц), dauer.ru (79, из них 25 старых адресов отдают 404), osnovit.ru (640; к нему добавлена карта сайта).',
    'Карточек товара: Perel 197, Dauer 45, Основит 247 (у Основита один товар лежит в нескольких разделах — склеено по canonical). Из них страниц-подборок без своего товара: Perel 7, Основит 13.',
    'Каждая карточка разобрана в запись: назначение, фасовка, расход, толщина, норматив, до 14 характеристик — ДОСЛОВНО. Все 13 465 цитат программно сверены с текстом карточки: несовпадений 0.',
    '«Товар» = продукт без цвета, фасовки и зимней модификации (Perel SL — один товар, 28 карточек: 14 цветов × 25/50 кг).',
    'Три карточки Основита (Флайвэлл PC24/1 ML, Стартвэлл PC21, PC21 M) не учтены: судя по разделу «Штукатурка», это старые названия серии «Техно».',
    'Цены: у Perel на сайте показаны, у Dauer и Основита нет. На наш сайт цены производителей не ставим.',
    'Внутренние противоречия у производителей (расход в тексте и в характеристиках разный, «мешки по 25 кг» на карточке 50 кг и т.п.) — в столбце «Примечание разбора» листа «Товары»; спорные цифры на сайт без уточнения не выводим.',
    'Не собраны: ссылки на техкарты Основита (на страницах названы, адреса не вытащились обходом) — добрать при сборке.',
]
for i, t in enumerate(notes, 1):
    c = wm.cell(row=i, column=1, value=t); c.font = Font(name=F, size=11 if i == 1 else 10, bold=(i == 1)); c.alignment = wrap
wm.column_dimensions['A'].width = 130
out = 'out/Сухие смеси — ассортимент трёх производителей (09.10.2026).xlsx'
wb.calculation.fullCalcOnLoad = True
wb.save(out); print(out, 'товаров', len(rows))
